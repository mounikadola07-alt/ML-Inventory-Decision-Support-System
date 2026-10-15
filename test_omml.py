import docx
from docx.oxml import parse_xml
import win32com.client

doc = docx.Document()

# Equation 1: EOQ = sqrt( (2*K*D) / h )
p1 = doc.add_paragraph()
omml_eoq = """
<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
    <m:oMath>
        <m:r><m:rPr><m:sty m:val="p"/></m:rPr><m:t>EOQ</m:t></m:r>
        <m:r><m:t> = </m:t></m:r>
        <m:rad>
            <m:radPr><m:degHide m:val="on"/></m:radPr>
            <m:deg/>
            <m:e>
                <m:f>
                    <m:num>
                        <m:r><m:t>2</m:t></m:r>
                        <m:r><m:rPr><m:scr m:val="roman"/></m:rPr><m:t>K</m:t></m:r>
                        <m:r><m:t>D̂</m:t></m:r>
                    </m:num>
                    <m:den>
                        <m:r><m:t>h</m:t></m:r>
                    </m:den>
                </m:f>
            </m:e>
        </m:rad>
    </m:oMath>
</m:oMathPara>
"""
p1._p.append(parse_xml(omml_eoq))

# Equation 2: SS = Z_alpha * sqrt( L_bar * sigma_e^2 + d_hat^2 * sigma_L^2 )
p2 = doc.add_paragraph()
omml_ss = """
<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
    <m:oMath>
        <m:r><m:rPr><m:sty m:val="p"/></m:rPr><m:t>SS</m:t></m:r>
        <m:r><m:t> = </m:t></m:r>
        <m:sSub>
            <m:e><m:r><m:t>Z</m:t></m:r></m:e>
            <m:sub><m:r><m:t>α</m:t></m:r></m:sub>
        </m:sSub>
        <m:rad>
            <m:radPr><m:degHide m:val="on"/></m:radPr>
            <m:deg/>
            <m:e>
                <m:r><m:t>L̄</m:t></m:r>
                <m:sSubSup>
                    <m:e><m:r><m:t>σ</m:t></m:r></m:e>
                    <m:sub><m:r><m:t>e</m:t></m:r></m:sub>
                    <m:sup><m:r><m:t>2</m:t></m:r></m:sup>
                </m:sSubSup>
                <m:r><m:t> + </m:t></m:r>
                <m:sSup>
                    <m:e><m:r><m:t>d̂</m:t></m:r></m:e>
                    <m:sup><m:r><m:t>2</m:t></m:r></m:sup>
                </m:sSup>
                <m:sSubSup>
                    <m:e><m:r><m:t>σ</m:t></m:r></m:e>
                    <m:sub><m:r><m:t>L</m:t></m:r></m:sub>
                    <m:sup><m:r><m:t>2</m:t></m:r></m:sup>
                </m:sSubSup>
            </m:e>
        </m:rad>
    </m:oMath>
</m:oMathPara>
"""
p2._p.append(parse_xml(omml_ss))

doc.save('test_math.docx')
print('Complex OMML test saved successfully!')

word = win32com.client.Dispatch('Word.Application')
word.Visible = False
try:
    d = word.Documents.Open(r'D:\College\Second Year\Second year 1st sem\ML\cat\test_math.docx')
    print('Recognized OMath count:', d.OMaths.Count)
    for i in range(1, d.OMaths.Count + 1):
        print(f'Math {i} Range Text: {repr(d.OMaths.Item(i).Range.Text.strip())}')
    d.Close(False)
finally:
    word.Quit()
