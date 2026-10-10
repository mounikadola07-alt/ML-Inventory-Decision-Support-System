"""
build_cover_letter.py - Generates a formal Journal Submission Cover Letter
in both Word (.docx) and Markdown (.md) formats.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_cover_letter():
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)
    
    def add_p(text="", bold=False, italic=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        if text:
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.bold = bold
            run.italic = italic
        return p

    # Sender Info
    add_p("Department of Computer Science and Engineering\nKoneru Lakshmaiah Education Foundation\nHyderabad, Telangana, India\nEmail: sashank.p@klh.edu.in", size=10, space_after=12)
    
    # Date
    add_p("10 October 2026", space_after=12)
    
    # Recipient Info
    add_p("To the Editor-in-Chief\nInternational Journal of Forecasting\nElsevier Science B.V.", bold=True, space_after=12)
    
    # Subject
    add_p("Subject: Submission of Original Empirical Research Article – 'An Empirically Audited Machine Learning and Stochastic Inventory Control Framework for Enterprise Replenishment Decision Support'", bold=True, space_after=10)
    
    # Salutation
    add_p("Dear Editor-in-Chief and Editorial Board Members,")
    
    # Body Paragraph 1
    add_p(
        "We are pleased to submit our original research manuscript entitled 'An Empirically Audited Machine Learning and Stochastic Inventory Control Framework for Enterprise Replenishment Decision Support' for publication consideration in the International Journal of Forecasting as a Regular Research Paper."
    )
    
    # Body Paragraph 2: Operational Motivation & Problem
    add_p(
        "Modern enterprise supply chains operate under an asymmetric cost structure where customer stockouts inflict catastrophic lost margins, contractual service-level penalties, and permanent customer defection, whereas holding excess safety buffer stock incurs marginal carrying charges (~20%/year). Despite this fundamental economic asymmetry, prevailing machine learning forecasting deployments routinely select algorithms using symmetric point loss objectives (MAE and RMSE). We investigate the critical operational question: Why do models that win statistical accuracy competitions frequently trigger catastrophic stockouts when deployed inside automated enterprise replenishment systems?"
    )
    
    # Body Paragraph 3: Theoretical & Methodological Contributions
    add_p(
        "Our investigation resolves this disconnect through three core scientific contributions:\n"
        "1. Analytical Proofs of Negative Bias Traps: We prove via Jensen's Inequality that fitting log-transformed targets (Log1p) forces predictions to the conditional geometric median rather than the conditional mean, inducing an extreme -55.9% to -81.3% under-forecasting bias. Furthermore, we prove that Tweedie compound Poisson-Gamma deviance exhibits an explosive gradient divergence at zero demand (lim_{μ→0+} ∂d/∂μ = +∞), exerting an overpowering downward force that drives chronic replenishment starvation.\n"
        "2. Tri-Criterion Operational Model Selection Rule: We establish a robust selection protocol requiring models to pass a strict operational bias constraint (|Bias %| ≤ 5.0%) prior to error ranking, preventing downstream inventory starvation without sacrificing point forecasting fidelity.\n"
        "3. Capital Optimization via Eve's Law of Total Variance: Grounded in stochastic inventory theory, we decouple deterministic predictable variance from stochastic risk, sizing safety stocks directly from forecast residual standard errors (σ_e) across lead-time demand convolution. In closed-loop simulations across 8,000 demand interactions, our proposed framework eliminated 24 stockout days, delivered a 92.9% reduction in total inventory operating costs, and liberated 49.9% of safety buffer working capital."
    )
    
    # Body Paragraph 4: Novel Surge Classifier
    add_p(
        "Furthermore, we formulate an integrated Platt-calibrated demand surge early-warning classifier, achieving an out-of-sample ROC AUC of 0.9278 with 83.21% classification accuracy across heterogeneous retail regimes, empowering supply chain controllers to execute proactive buffer replenishment before promotional surges deplete on-hand stock."
    )
    
    # Body Paragraph 5: Reproducibility & Declarations
    add_p(
        "To ensure uncompromising scientific transparency and reproducibility, all benchmark datasets, feature engineering pipelines, walk-forward cross-validation routines, and simulation engines have been audited, packaged with 90 passing automated unit/regression tests, and made accessible as an interactive cloud notebook runnable in Google Colab."
    )
    add_p(
        "We confirm that this manuscript represents original work that has not been published previously and is not currently under consideration for publication elsewhere. All authors have approved the manuscript and agreed to its submission."
    )
    add_p(
        "Thank you for your time and editorial consideration. We look forward to hearing from you."
    )
    
    # Sign-off
    add_p("Sincerely,\n\nSashank P. (Corresponding Author)\nOn behalf of all co-authors: Abhinay Sai, Poli Naidu, and Sushama Rani Dutta\nDepartment of Computer Science and Engineering\nKoneru Lakshmaiah Education Foundation, Hyderabad, Telangana, India", space_before=12)

    docx_path = "Journal_Submission_Cover_Letter.docx"
    doc.save(docx_path)
    print(f"Successfully generated {docx_path} ({os.path.getsize(docx_path)} bytes)!")

    # Markdown export
    md_content = r"""# Journal Submission Cover Letter

**Department of Computer Science and Engineering**  
Koneru Lakshmaiah Education Foundation, Hyderabad, Telangana, India  
Email: `sashank.p@klh.edu.in`  

**Date:** 10 October 2026  

**To:**  
The Editor-in-Chief  
*International Journal of Forecasting*  
Elsevier Science B.V.  

**Subject:** Submission of Original Empirical Research Article – *"An Empirically Audited Machine Learning and Stochastic Inventory Control Framework for Enterprise Replenishment Decision Support"*  

Dear Editor-in-Chief and Editorial Board Members,

We are pleased to submit our original research manuscript entitled **"An Empirically Audited Machine Learning and Stochastic Inventory Control Framework for Enterprise Replenishment Decision Support"** for publication consideration in the *International Journal of Forecasting* as a Regular Research Paper.

Modern enterprise supply chains operate under an asymmetric cost structure where customer stockouts inflict catastrophic lost margins, contractual service-level penalties, and permanent customer defection, whereas holding excess safety buffer stock incurs marginal carrying charges (~20%/year). Despite this fundamental economic asymmetry, prevailing machine learning forecasting deployments routinely select algorithms using symmetric point loss objectives (MAE and RMSE). We investigate the critical operational question: *Why do models that win statistical accuracy competitions frequently trigger catastrophic stockouts when deployed inside automated enterprise replenishment systems?*

Our investigation resolves this disconnect through three core scientific contributions:
1. **Analytical Proofs of Negative Bias Traps:** We prove via Jensen's Inequality that fitting log-transformed targets (Log1p) forces predictions to the conditional geometric median rather than the conditional mean, inducing an extreme -55.9% to -81.3% under-forecasting bias. Furthermore, we prove that Tweedie compound Poisson-Gamma deviance exhibits an explosive gradient divergence at zero demand ($\lim_{\mu \to 0^+} \frac{\partial d(0, \mu)}{\partial \mu} = +\infty$), exerting an overpowering downward force that drives chronic replenishment starvation.
2. **Tri-Criterion Operational Model Selection Rule:** We establish a robust selection protocol requiring models to pass a strict operational bias constraint ($|\text{Bias}\%| \le 5.0\%$) prior to error ranking, preventing downstream inventory starvation without sacrificing point forecasting fidelity.
3. **Capital Optimization via Eve's Law of Total Variance:** Grounded in stochastic inventory theory, we decouple deterministic predictable variance from stochastic risk, sizing safety stocks directly from forecast residual standard errors ($\sigma_e$) across lead-time demand convolution. In closed-loop simulations across 8,000 demand interactions, our proposed framework eliminated 24 stockout days, delivered a 92.9% reduction in total inventory operating costs, and liberated 49.9% of safety buffer working capital.

Furthermore, we formulate an integrated Platt-calibrated demand surge early-warning classifier, achieving an out-of-sample **ROC AUC of 0.9278** with **83.21% classification accuracy** across heterogeneous retail regimes, empowering supply chain controllers to execute proactive buffer replenishment before promotional surges deplete on-hand stock.

To ensure uncompromising scientific transparency and reproducibility, all benchmark datasets, feature engineering pipelines, walk-forward cross-validation routines, and simulation engines have been audited, packaged with 90 passing automated unit/regression tests, and made accessible as an interactive cloud notebook runnable in Google Colab.

We confirm that this manuscript represents original work that has not been published previously and is not currently under consideration for publication elsewhere. All authors have approved the manuscript and agreed to its submission.

Thank you for your time and editorial consideration. We look forward to hearing from you.

Sincerely,  

**Sashank P.** *(Corresponding Author)*  
On behalf of co-authors: Abhinay Sai, Poli Naidu, and Sushama Rani Dutta  
Department of Computer Science and Engineering  
Koneru Lakshmaiah Education Foundation, Hyderabad, Telangana, India  
"""
    with open("Journal_Submission_Cover_Letter.md", "w", encoding="utf-8") as f:
        f.write(md_content)
    print("Successfully generated Journal_Submission_Cover_Letter.md!")

if __name__ == "__main__":
    create_cover_letter()
