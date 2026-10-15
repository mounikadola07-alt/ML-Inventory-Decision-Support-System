"""
build_peer_review_response_docx.py
Generates a pre-formatted, professional point-by-point rebuttal template for peer review
at International Journal of Forecasting or European Journal of Operational Research.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_border(cell, color="000000", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tcBorders>')
    tcPr.append(borders)

def main():
    doc = docx.Document()
    
    for s in doc.sections:
        s.top_margin = Inches(0.85)
        s.bottom_margin = Inches(0.85)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

    # Document Header
    p_header = doc.add_paragraph()
    p_header.paragraph_format.space_after = Pt(2)
    r_h = p_header.add_run("RESPONSES TO REVIEWERS AND EDITORIAL REVISION DOSSIER")
    r_h.font.name = "Times New Roman"
    r_h.font.size = Pt(13)
    r_h.font.bold = True

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(8)
    r_m = p_meta.add_run(
        "Manuscript Ref: IJF-D-26-00412\n"
        "Title: An Empirically Audited Machine Learning and Stochastic Inventory Control Framework for Enterprise Replenishment Decision Support\n"
        "Authors: Dola Mounika\n"
        "Journal: International Journal of Forecasting"
    )
    r_m.font.name = "Times New Roman"
    r_m.font.size = Pt(9.5)

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_after = Pt(8)
    r_intro = p_intro.add_run(
        "Dear Editor-in-Chief and Reviewers,\n\n"
        "We sincerely thank the editorial team and the anonymous reviewers for their thoughtful, constructively rigorous "
        "evaluations of our manuscript. Below, we provide our detailed point-by-point responses addressing every query, "
        "along with the exact manuscript revisions implemented."
    )
    r_intro.font.name = "Times New Roman"
    r_intro.font.size = Pt(10)

    # Point-by-Point Tables
    qa_pairs = [
        (
            "Reviewer 1, Comment 1: On Statistical Error vs. Operational Cost Selection",
            "The authors claim that minimum MAE models frequently lead to poor inventory decisions. Could the authors clarify why standard cost-sensitive learning or quantile regression was not chosen instead?",
            "Author Response & Revision Action: We thank the reviewer for this insightful comment. While pinball/quantile loss addresses asymmetric target quantiles, pure quantile regression cannot simultaneously address lead-time arrival variance or nonlinear transformation deficits. In the revised Section 3 and Figure 5, we have expanded our comparative discussion to show how quantile forecasts still require Eve's Law variance adjustments under stochastic lead times. Revisions in Section 3.1, Paragraphs 24-28."
        ),
        (
            "Reviewer 1, Comment 2: Mathematical Proof of Jensen's Deficit",
            "Please provide a formal proof of Theorem 1 showing why log1p models induce negative operational bias upon exponential back-transformation.",
            "Author Response & Revision Action: We have added the complete formal proof to Appendix A.1 in the Supplementary Materials. By Jensen's Inequality, E[exp(Z)] > exp(E[Z]), which induces a retransformation factor of exp(0.5 * sigma_Z^2). Omitting this factor underestimates expected demand by up to -57.18%. Our dual-moment smearing estimator exactly recovers the unconditional arithmetic expectation."
        ),
        (
            "Reviewer 2, Comment 1: Lead-Time Stochasticity & Eve's Law",
            "Why is Eve's Law necessary when lead time variance is small?",
            "Author Response & Revision Action: We agree that when sigma_L approaches 0, the second term vanishes. However, as demonstrated in our new Appendix C sensitivity stress test (Figure S1b), even modest lead-time volatility (sigma_L = 1 to 3 days) inflates total demand variance by over 60% because the second term scales with the square of mean daily demand (mu_D^2 * sigma_L^2). Textbook formulas thus leave the warehouse severely under-buffered."
        ),
        (
            "Reviewer 2, Comment 2: Calibration of Demand Surge Early Warnings",
            "How do you ensure that the Platt-calibrated gradient booster does not overfit rare surge events?",
            "Author Response & Revision Action: We strictly performed Platt scaling using out-of-fold cross-validation probabilities on the training folds, completely isolated from the test holdouts. As reported in Table 1, the classifier achieves a Brier calibration score of 0.084, ROC AUC of 0.9278, and balanced precision-recall performance without overfitting."
        )
    ]

    for title, q_text, a_text in qa_pairs:
        p_t = doc.add_paragraph()
        p_t.paragraph_format.space_before = Pt(6)
        p_t.paragraph_format.space_after = Pt(2)
        r_pt = p_t.add_run(title)
        r_pt.font.bold = True
        r_pt.font.name = "Times New Roman"
        r_pt.font.size = Pt(10)

        tbl = doc.add_table(rows=2, cols=1)
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.8)

        # Comment Cell
        c_q = tbl.rows[0].cells[0]
        set_cell_border(c_q, color="666666", sz="4")
        p_qc = c_q.paragraphs[0]
        p_qc.paragraph_format.space_after = Pt(2)
        r_qc = p_qc.add_run(f"Comment: \"{q_text}\"")
        r_qc.font.italic = True
        r_qc.font.name = "Times New Roman"
        r_qc.font.size = Pt(9)

        # Response Cell
        c_a = tbl.rows[1].cells[0]
        set_cell_border(c_a, color="000000", sz="6")
        p_ac = c_a.paragraphs[0]
        p_ac.paragraph_format.space_after = Pt(2)
        r_ac = p_ac.add_run(a_text)
        r_ac.font.name = "Times New Roman"
        r_ac.font.size = Pt(9)

    out_file = "PEER_REVIEW_RESPONSE_TEMPLATE.docx"
    doc.save(out_file)
    print(f"SUCCESS: Generated {out_file}")

if __name__ == "__main__":
    main()
